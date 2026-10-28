"""Сборка руководств Burndown из Markdown в PDF и DOCX.

    python3 tools/docs/build_manuals.py              # все руководства
    python3 tools/docs/build_manuals.py user-guide   # одно

Нужны pandoc 3.x, python-docx, pypdf, Playwright с Chromium и LibreOffice (оглавление DOCX).
Источники — docs/manuals/*.md и docs/manuals/images (снимки: tools/docs/shots.py); результат кладётся рядом с .md.
PDF собирается в два прохода: после первого номера страниц разделов читаются из ссылок содержания
и вписываются в него во втором.
"""
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from playwright.sync_api import sync_playwright
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[2]
MANUALS = ROOT / 'docs' / 'manuals'
TOOLS = Path(__file__).resolve().parent
VERSION = '1.1.0'
DOCS = {
    'user-guide': 'Руководство пользователя',
    'admin-guide': 'Руководство администратора',
}
TEAL = RGBColor(0x0C, 0x4E, 0x54)
INK = RGBColor(0x1C, 0x21, 0x26)
MUTED = RGBColor(0x6B, 0x73, 0x85)
FONT = 'Arial'


def run(cmd, **kw):
    subprocess.run(cmd, check=True, **kw)


def footer_text(title):
    return f'Burndown {VERSION} · {title}'


# ---------------------------------------------------------------- PDF
def toc_page_numbers(pdf_path):
    """Номера страниц, на которые ведут ссылки содержания (в порядке ссылок)."""
    reader = PdfReader(str(pdf_path))
    page_index = {page.indirect_reference.idnum: i for i, page in enumerate(reader.pages)}
    named = {}
    try:
        for key, dest in reader.named_destinations.items():
            named[key] = reader.get_destination_page_number(dest)
    except Exception:  # noqa: BLE001 — разные версии Chromium пишут назначения по-разному
        pass
    numbers = []
    for page in reader.pages[1:6]:  # содержание сразу после обложки
        links = []
        for ref in page.get('/Annots') or []:
            annot = ref.get_object()
            if annot.get('/Subtype') != '/Link':
                continue
            dest = annot.get('/Dest')
            if dest is None and annot.get('/A'):
                dest = annot['/A'].get('/D')
            target = None
            if isinstance(dest, str) or hasattr(dest, 'startswith'):
                target = named.get(str(dest))
            elif dest is not None:
                target = page_index.get(dest[0].idnum)
            if target is None:
                continue
            rect = annot.get('/Rect')
            links.append((-float(rect[3]), float(rect[0]), target + 1))
        if not links and numbers:
            break
        numbers.extend(n for *_pos, n in sorted(links))
    return numbers


def render_pdf(browser, html, out, footer, numbers=None, cover_only=False):
    """HTML → PDF в Chromium: A4, колонтитул с номером страницы, закладки по заголовкам."""
    page = browser.new_page()
    page.goto(html.as_uri(), wait_until='load')
    page.evaluate('document.fonts.ready')
    # номера страниц в содержании; в первом проходе — резерв места, чтобы вёрстка не сдвинулась
    page.evaluate("""(nums) => document.querySelectorAll('nav.toc a').forEach((a, i) => {
        const span = document.createElement('span'); span.className = 'toc-page';
        span.textContent = nums ? (nums[i] == null ? '' : String(nums[i])) : '00'; a.appendChild(span);
    })""", numbers)
    page.pdf(path=str(out), format='A4', print_background=True, outline=True, tagged=True,
             prefer_css_page_size=True, display_header_footer=not cover_only,
             page_ranges='1' if cover_only else '', header_template='<span></span>',
             footer_template=f'''<div style="width:100%;font-size:7.5px;font-family:Inter,Arial,sans-serif;color:#8a92a6;
               padding:0 17mm;display:flex;justify-content:space-between"><span>{footer}</span>
               <span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>''')
    page.close()


def replace_cover(pdf, cover):
    """Первая страница — обложка без колонтитула; закладки и ссылки остальных страниц сохраняются."""
    writer = PdfWriter(clone_from=str(pdf))
    writer.insert_page(PdfReader(str(cover)).pages[0], 0)
    writer.remove_page(1)
    writer.write(str(pdf))


def build_pdf(browser, md, title):
    html = md.with_suffix('.build.html')
    css = (TOOLS / 'manual.css').read_text(encoding='utf-8')
    run(['pandoc', str(md), '-f', 'markdown+implicit_figures+link_attributes+ascii_identifiers', '-t', 'html5',
         '--template', str(TOOLS / 'template.html'), '--toc', '--toc-depth=2', '--number-sections',
         '-V', f'css={css}', '-V', f'cover-title={title}', '--resource-path', str(md.parent), '-o', str(html)])
    pdf = md.with_suffix('.pdf')
    footer = footer_text(title)
    try:
        render_pdf(browser, html, pdf, footer)
        numbers = toc_page_numbers(pdf)
        render_pdf(browser, html, pdf, footer, numbers)
        cover = pdf.with_suffix('.cover.pdf')
        render_pdf(browser, html, cover, footer, numbers, cover_only=True)
        replace_cover(pdf, cover)
        cover.unlink()
    finally:
        html.unlink(missing_ok=True)
    print(f'  pdf: {pdf.relative_to(ROOT)} ({len(PdfReader(str(pdf)).pages)} стр., ссылок содержания: {len(numbers)})')


# ---------------------------------------------------------------- DOCX
def reference_docx(path):
    data = subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'], check=True,
                          capture_output=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    styles = {s.name: s for s in doc.styles}

    def font(style, size=None, bold=None, color=None, name=FONT):
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rfonts.set(qn(attr), name)
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            rfonts.attrib.pop(qn(attr), None)
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color is not None:
            style.font.color.rgb = color

    for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        if name in styles:
            font(styles[name], 10.5, color=INK)
            styles[name].paragraph_format.space_after = Pt(5)
            styles[name].paragraph_format.line_spacing = 1.15
    font(styles['Title'], 28, True, INK)
    font(styles['Subtitle'], 15, False, TEAL)
    for name in ('Author', 'Date'):
        if name in styles:
            font(styles[name], 11, False, MUTED)
    for level, size in {1: 18, 2: 14, 3: 12, 4: 11}.items():
        style = styles[f'Heading {level}']
        font(style, size, True, TEAL if level == 2 else INK)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles['Heading 1'].paragraph_format.page_break_before = True
    for name in ('Image Caption', 'Caption', 'Table Caption'):
        if name in styles:
            font(styles[name], 9, False, MUTED)
            styles[name].font.italic = False
    if 'Source Code' in styles:
        font(styles['Source Code'], 8.5, name='Consolas')
    if 'Block Text' in styles:
        font(styles['Block Text'], 10, False, INK)
    if 'TOC Heading' in styles:
        font(styles['TOC Heading'], 18, True, INK)
    doc.save(str(path))


def shade(element, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    element.append(shd)


def polish_docx(path, title):
    doc = Document(str(path))
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:color'), 'D3DCDA')
            borders.append(b)
        table._tbl.tblPr.append(borders)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                if i == 0:
                    shade(cell._tc.get_or_add_tcPr(), 'E6F0EF')
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for r in paragraph.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
    for paragraph in doc.paragraphs:  # выноски — мягкая заливка
        if paragraph.style.name == 'Block Text':
            shade(paragraph._p.get_or_add_pPr(), 'EEF6F5')
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = footer_text(title) + ' · стр. '
    page_run = footer.add_run()
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            fld = OxmlElement('w:fldChar')
            fld.set(qn('w:fldCharType'), kind)
            page_run._r.append(fld)
        else:
            instr = OxmlElement('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = text
            page_run._r.append(instr)
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x8A, 0x92, 0xA6)
    doc.save(str(path))


def build_docx(md, title, reference):
    out = md.with_suffix('.docx')
    run(['pandoc', str(md), '-f', 'markdown+implicit_figures+link_attributes', '-t', 'docx',
         '--reference-doc', str(reference), '--toc', '--toc-depth=2', '--number-sections',
         '-M', 'toc-title=Содержание', '--resource-path', str(md.parent), '-o', str(out)])
    polish_docx(out, title)
    # оглавление с номерами страниц: LibreOffice обновляет поле TOC и сохраняет DOCX
    run([sys.executable, str(TOOLS / 'update_docx_toc.py'), str(out)])
    print(f'  docx: {out.relative_to(ROOT)}')


def main(names):
    with tempfile.TemporaryDirectory() as tmp, sync_playwright() as p:
        reference = Path(tmp) / 'reference.docx'
        reference_docx(reference)
        browser = p.chromium.launch(executable_path=os.environ.get('CHROMIUM_PATH') or None)
        for name in names:
            title = DOCS[name]
            md = MANUALS / f'{name}.md'
            print(name)
            build_docx(md, title, reference)
            build_pdf(browser, md, title)
        browser.close()


if __name__ == '__main__':
    main(sys.argv[1:] or [n for n in DOCS if (MANUALS / f'{n}.md').exists()])
